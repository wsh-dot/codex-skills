"""Read local EC5 project packages without the application; author source snapshots with artifact-tool."""
from pathlib import Path
import argparse, hashlib, json, math, os, re, shutil, struct, subprocess, sys, time, uuid, zipfile
from ecpt5_model import build_model

PKG=Path(__file__).resolve().parent.parent
JAR=PKG/'assets/db4o-6.4.38.10595-java5.jar'

def run(cmd, timeout=60):
    try:
        p=subprocess.run([str(x) for x in cmd],capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=timeout)
    except subprocess.TimeoutExpired as exc:raise RuntimeError(f'PROCESS_TIMEOUT after {timeout}s: {Path(str(cmd[0])).name}') from exc
    if p.returncode:raise RuntimeError((p.stderr or p.stdout or f'exit={p.returncode}')[-1800:])
    return p

def preflight_database(raw):
    offset=raw[:64].find(b'db4o')
    if offset<0:raise ValueError('DB_SIGNATURE_MISSING')
    b=raw[offset:]
    if len(b)<100 or b[4]!=1:raise ValueError('DB_HEADER_UNSUPPORTED')
    bs=b[36]
    if not 1<=bs<=127:raise ValueError('DB_BLOCK_SIZE_INVALID')
    variable=struct.unpack_from('>i',b,37)[0]*bs
    if not 0<=variable<=len(b)-16:raise ValueError('DB_HEADER_POINTER_INVALID')
    address,length=struct.unpack_from('>ii',b,variable)
    position=address*bs
    if not 0<=position<=len(b)-4:raise ValueError('DB_REPOSITORY_POINTER_INVALID')
    count=struct.unpack_from('>i',b,position)[0]
    if not 0<count<=100000 or length<4+4*count or position+length>len(b):
        raise ValueError(f'DB_REPOSITORY_INVALID offset={position} count={count} length={length}')
    return b,{'prefixBytes':offset,'blockSize':bs,'repositoryOffset':position,'classCount':count}

def extract_project(source,work,java,classes,timeout):
    failures={};headers={}
    with zipfile.ZipFile(source) as z:
        names=z.namelist()
        if len(names)!=len(set(names)):raise ValueError('DUPLICATE_ZIP_ENTRY')
        if 'arr' not in names:raise ValueError('PROJECT_ARR_MISSING')
        entries=[i for i in z.infolist() if i.filename.endswith('.ecd')]
        if not entries:raise ValueError('NO_ENGINEERING_DATABASES')
        if sum(i.file_size for i in z.infolist())>1024**3:raise ValueError('PACKAGE_EXCEEDS_1GIB_INSPECT_BEFORE_EXPORT')
        bad=z.testzip()
        if bad:raise ValueError('ZIP_CRC_FAILED: '+bad)
        for entry in entries:
            if not re.fullmatch(r'[A-Za-z0-9_-]+\.ecd',entry.filename):raise ValueError('UNSUPPORTED_DATABASE_ENTRY: '+entry.filename)
            uid=entry.filename[:-4]
            raw=z.read(entry)
            try:
                data,header=preflight_database(raw);headers[uid]=header
                target=work/entry.filename;target.write_bytes(data)
                temp=work/(entry.filename+'.jsonl.pending')
                run([java,'--add-opens','java.base/java.util=ALL-UNNAMED','--add-opens','java.base/java.lang=ALL-UNNAMED','-Xmx1g','-cp',os.pathsep.join([str(classes),str(JAR)]),'Ecpt5Reader','dump',target,temp],timeout)
                temp.replace(work/(entry.filename+'.jsonl'))
            except Exception as exc:failures[uid]=str(exc)
    run([java,'-cp',os.pathsep.join([str(classes),str(JAR)]),'Ecpt5Reader','arr',source,work/'arr.xml'],timeout)
    return failures,headers

def verify_workbook(file,d):
    import openpyxl
    wb=openpyxl.load_workbook(file,data_only=False,read_only=True)
    values=openpyxl.load_workbook(file,data_only=True,read_only=True)
    checked=0
    for name,key in [('树结构','rows'),('定额人材机','spends'),('项目其他费用','fees'),('来源与核对','summary')]:
        actual=list(wb[name].iter_rows(min_row=5,values_only=True));expected=d[key]
        if len(actual)!=len(expected):raise ValueError(f'WORKBOOK_ROW_COUNT_MISMATCH: {name}')
        for i,(a,b) in enumerate(zip(actual,expected),5):
            for j,(x,y) in enumerate(zip(a,b),1):
                if isinstance(x,(float,int)) and isinstance(y,(float,int)):same=math.isclose(x,y,rel_tol=1e-12,abs_tol=1e-8)
                else:same=x==y or (x in (None,'') and y in (None,''))
                if not same:raise ValueError(f'WORKBOOK_VALUE_MISMATCH: {name}!R{i}C{j}: {x!r} != {y!r}')
                checked+=1
    formula_checks=0
    for i,row in enumerate(d['checks'],5):
        f=wb['工程核对'].cell(i,10).value;v=values['工程核对'].cell(i,10).value
        if row[8] is None:
            if f is not None or v is not None:raise ValueError('UNKNOWN_AMOUNT_CONVERTED_TO_ZERO')
        else:
            if f!=f'=I{i}-H{i}' or not isinstance(v,(int,float)) or abs(v-row[9])>1e-8:raise ValueError('RECONCILIATION_FORMULA_FAILED')
            formula_checks+=1
    errors=[(s.title,c.coordinate,c.value) for s in wb for row in s for c in row if c.data_type=='e']
    if errors:raise ValueError('WORKBOOK_ERROR_CELLS: '+str(errors[:5]))
    wb.close();values.close()
    return {'sourceValuesChecked':checked,'formulaCachesChecked':formula_checks,'errorCells':0}

def link_modules(folder,target):
    link=folder/'node_modules'
    if link.exists():raise ValueError('NODE_MODULES_LINK_ALREADY_EXISTS')
    if os.name=='nt':run(['cmd','/c','mklink','/J',str(link),str(target)])
    else:link.symlink_to(target,target_is_directory=True)

class ExportParser(argparse.ArgumentParser):
    def error(self,message):self.exit(1,f'Configuration error: {message}\n')

def main(argv=None):
    p=ExportParser(description='Export EC5 .ecpt5 project files to independent Excel workbooks. Exit 0=complete, 2=partial, 1=failed.')
    p.add_argument('inputs',nargs='+',type=Path)
    p.add_argument('--output-dir',type=Path,required=True)
    p.add_argument('--node-exe',type=Path)
    p.add_argument('--node-modules',type=Path)
    p.add_argument('--java',default=shutil.which('java'))
    p.add_argument('--javac',default=shutil.which('javac'))
    p.add_argument('--database-timeout',type=int,default=45)
    p.add_argument('--extract-only',action='store_true',help='Produce verified data.json and manifests without authoring Excel.')
    args=p.parse_args(argv)
    if not args.java or not args.javac:p.error('Java JDK 17+ is required; specify --java and --javac.')
    if args.database_timeout<1:p.error('--database-timeout must be positive.')
    if not args.extract_only and (not args.node_exe or not args.node_modules):p.error('Load workspace dependencies, then supply --node-exe and --node-modules.')
    out=args.output_dir.resolve();out.mkdir(parents=True,exist_ok=True)
    workroot=out/('.ecpt5-work-'+uuid.uuid4().hex[:12]);workroot.mkdir()
    classes=workroot/'classes';classes.mkdir()
    run([args.javac,'-encoding','UTF-8','-cp',JAR,'-d',classes,PKG/'scripts/Ecpt5Reader.java'])
    if not args.extract_only:
        modules=args.node_modules.resolve()
        if not (modules/'@oai/artifact-tool/package.json').is_file():p.error('@oai/artifact-tool is unavailable at --node-modules.')
        link_modules(workroot,modules)
        shutil.copy2(PKG/'scripts/build_workbook.mjs',workroot/'build_workbook.mjs')
    results=[]
    for i,path in enumerate(args.inputs):
        src=path.resolve();item={'source':str(src)}
        try:
            if src.suffix.lower()!='.ecpt5':raise ValueError('EXPECTED_ECPT5')
            before=hashlib.sha256(src.read_bytes()).hexdigest();item['sourceSha256']=before
            work=workroot/f'{i:03d}-{before[:10]}';work.mkdir()
            failures,headers=extract_project(src,work,args.java,classes,args.database_timeout)
            label=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',src.stem)[:55]+'_'+before[:8]
            d=build_model(work,src,failures,label)
            item.update({'status':'complete' if d['isComplete'] else 'partial','expectedDatabases':d['expectedDatabases'],'readDatabases':d['readDatabases'],'treeRows':len(d['rows']),'resourceRows':len(d['spends']),'otherFeeRows':len(d['fees']),'unreadable':d['unreadable'],'workDir':str(work),'headerChecks':headers,'amountDifferences':[[c[0],c[2],c[9]] for c in d['checks'] if c[9] is not None and abs(c[9])>1e-5]})
            if not args.extract_only:
                target=out/d['outputName']
                if target.exists():raise ValueError('OUTPUT_ALREADY_EXISTS: '+str(target))
                (work/'data.json').write_text(json.dumps(d,ensure_ascii=False,allow_nan=False),encoding='utf-8')
                # Builder writes beside work/; publish only after verification.
                cmd=[str(args.node_exe),str(workroot/'build_workbook.mjs'),str(work)]
                proc=subprocess.run(cmd,capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=300)
                (work/'authoring.log').write_text(proc.stdout+'\n'+proc.stderr,encoding='utf-8')
                saved=work.parent/d['outputName']
                if not (work/'authoring_complete.json').exists() or not saved.exists():raise RuntimeError('WORKBOOK_AUTHORING_FAILED: '+(proc.stderr or proc.stdout)[-1500:])
                item['verification']=verify_workbook(saved,d)
                saved.replace(target);item['xlsx']=str(target)
                item['authoringExitCode']=proc.returncode
            else:(work/'data.json').write_text(json.dumps(d,ensure_ascii=False,allow_nan=False),encoding='utf-8')
            if hashlib.sha256(src.read_bytes()).hexdigest()!=before:raise ValueError('SOURCE_CHANGED_DURING_EXPORT')
        except Exception as exc:item.update({'status':'failed','error':str(exc)})
        results.append(item)
        print(json.dumps({k:v for k,v in item.items() if k not in ('headerChecks','workDir')},ensure_ascii=False),flush=True)
    manifest=out/('manifest-'+workroot.name.removeprefix('.ecpt5-work-')+'.json')
    manifest.write_text(json.dumps({'results':results},ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
    print(json.dumps({'manifest':str(manifest),'complete':sum(x['status']=='complete' for x in results),'partial':sum(x['status']=='partial' for x in results),'failed':sum(x['status']=='failed' for x in results)},ensure_ascii=False))
    return 1 if any(x['status']=='failed' for x in results) else 2 if any(x['status']=='partial' for x in results) else 0

if __name__=='__main__':raise SystemExit(main())
